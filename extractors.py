"""Text extraction for PDF, Word, images, plain text, Markdown and CSV.

Each extractor returns an `Extraction` made of `Section`s. A section is a
logical unit with a human-readable location (page, heading, row range) that
is later split into passages by `chunking.py`.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from . import config as C


@dataclass
class Section:
    location: str
    text: str
    page: Optional[int] = None
    heading: str = ""


@dataclass
class Extraction:
    kind: str
    sections: list
    warnings: list = field(default_factory=list)
    pages: Optional[int] = None
    ocr_pages: int = 0
    title: str = ""


Transcriber = Callable[[object], str]   # bytes or PIL.Image -> text


# ── Helpers ─────────────────────────────────────────────────────────────
def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"([a-z])-\n([a-z])", r"\1\2", text)          # de-hyphenate line breaks
    text = re.sub(r"[ \t\u00a0\u200b]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _inline(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def compact_ranges(numbers) -> str:
    """[1,2,3,7,9,10] -> '1–3, 7, 9–10'"""
    numbers = sorted(set(numbers))
    out, start, prev = [], None, None
    for n in numbers:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            out.append(f"{start}–{prev}" if start != prev else str(start))
            start = prev = n
    if start is not None:
        out.append(f"{start}–{prev}" if start != prev else str(start))
    return ", ".join(out)


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# ── PDF ─────────────────────────────────────────────────────────────────
def _largest_page_image(page):
    best = None
    try:
        for img in page.images:
            size = len(img.data or b"")
            if best is None or size > best[0]:
                best = (size, img)
    except Exception:
        return None
    if not best:
        return None
    try:
        return best[1].image          # PIL image (handles JBIG2/CCITT via pypdf decoders)
    except Exception:
        return best[1].data


def extract_pdf(data: bytes, transcribe: Optional[Transcriber], ocr: bool) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise ValueError("This PDF is password-protected. Remove the password and upload it again.")
    total = len(reader.pages)
    if total > C.MAX_PDF_PAGES:
        raise ValueError(f"PDF has {total} pages; the maximum is {C.MAX_PDF_PAGES}.")

    title = ""
    try:
        title = _inline((reader.metadata or {}).get("/Title", "") or "")
    except Exception:
        pass

    sections, empty, ocr_failed, ocr_skipped = [], [], [], []
    ocr_done = 0
    for number, page in enumerate(reader.pages, 1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception:
            text = ""
        if len(text) >= 25:
            sections.append(Section(f"Page {number}", text, page=number))
            continue
        # Little or no text layer: probably a scanned page.
        if ocr and transcribe:
            if ocr_done >= C.MAX_OCR_PAGES_PER_PDF:
                ocr_skipped.append(number)
            else:
                image = _largest_page_image(page)
                if image is not None:
                    try:
                        ocr_text = clean_text(transcribe(image))
                        ocr_done += 1
                        if ocr_text:
                            sections.append(Section(f"Page {number} (OCR)", ocr_text, page=number))
                            continue
                    except ValueError:
                        pass                       # no legible text on the image
                    except Exception:
                        ocr_failed.append(number)
                        if text:
                            sections.append(Section(f"Page {number}", text, page=number))
                        continue
        if text:
            sections.append(Section(f"Page {number}", text, page=number))
        else:
            empty.append(number)

    warnings = []
    if ocr_done:
        warnings.append(f"{ocr_done} scanned page(s) were read by the vision model — verify important text.")
    if empty:
        hint = "" if ocr else " Turn on scanned-page reading to OCR them."
        warnings.append(f"No readable text on page(s) {compact_ranges(empty)}.{hint}")
    if ocr_failed:
        warnings.append(f"Vision reading failed on page(s) {compact_ranges(ocr_failed)}.")
    if ocr_skipped:
        warnings.append(f"OCR limit ({C.MAX_OCR_PAGES_PER_PDF} pages) reached; page(s) "
                        f"{compact_ranges(ocr_skipped)} were not read.")
    return Extraction("PDF", sections, warnings, pages=total, ocr_pages=ocr_done, title=title)


# ── Word ────────────────────────────────────────────────────────────────
def _heading_level(paragraph) -> Optional[int]:
    try:
        name = paragraph.style.name if paragraph.style is not None else ""
    except Exception:
        return None
    if name == "Title":
        return 0
    match = re.match(r"Heading (\d)", name or "")
    return int(match.group(1)) if match else None


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        seen, cells = [], []
        try:
            row_cells = row.cells
        except Exception:
            continue
        for cell in row_cells:
            if any(cell._tc is s for s in seen):      # merged cells repeat
                continue
            seen.append(cell._tc)
            cells.append(_inline(cell.text))
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_docx(data: bytes) -> Extraction:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > C.MAX_DOCX_UNCOMPRESSED_MB * 1024 * 1024:
                raise ValueError("Word file is too large after decompression.")
    except zipfile.BadZipFile:
        raise ValueError("Not a valid .docx file. Legacy .doc files must be saved as .docx first.")

    doc = Document(io.BytesIO(data))
    sections: list[Section] = []
    path: list[tuple[int, str]] = []          # heading stack
    buffer: list[str] = []
    first_para = para_no = table_no = 0

    def heading_path() -> str:
        return " › ".join(text for _, text in path[-2:])

    def flush():
        nonlocal buffer
        if buffer:
            span = f"¶{first_para}" if first_para == para_no else f"¶{first_para}–{para_no}"
            head = heading_path()
            sections.append(Section(f"{head} · {span}" if head else span, "\n\n".join(buffer), heading=head))
        buffer = []

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            level = _heading_level(paragraph)
            if level is not None:
                flush()
                while path and path[-1][0] >= level:
                    path.pop()
                path.append((level, _inline(text)[:80]))
                continue
            para_no += 1
            if not buffer:
                first_para = para_no
            buffer.append(text)
            if sum(len(b) for b in buffer) > 3000:     # keep sections a manageable size
                flush()
        elif tag == "tbl":
            flush()
            table_no += 1
            text = _table_text(Table(child, doc))
            if text.strip():
                head = heading_path()
                sections.append(Section(f"{head} · Table {table_no}" if head else f"Table {table_no}",
                                        text, heading=head))
    flush()

    # Text boxes and shapes (python-docx doesn't surface these via .paragraphs)
    boxes = []
    try:
        for box in doc.element.body.xpath(".//w:txbxContent"):
            text = "\n".join(_inline("".join(p.xpath(".//w:t/text()"))) for p in box.xpath(".//w:p"))
            text = text.strip()
            if text and text not in boxes:
                boxes.append(text)
    except Exception:
        pass
    if boxes:
        sections.append(Section("Text boxes", "\n\n".join(boxes)))

    # Headers & footers
    header_lines = []
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            try:
                for p in part.paragraphs:
                    t = _inline(p.text)
                    if t and t not in header_lines:
                        header_lines.append(t)
            except Exception:
                pass
    if header_lines:
        sections.append(Section("Headers & footers", "\n".join(header_lines)))

    title = ""
    try:
        title = _inline(doc.core_properties.title or "")
    except Exception:
        pass
    warnings = ["Word locations use headings and paragraph numbers (¶), not page numbers. "
                "Images inside the document and tracked changes are not read."]
    return Extraction("Word", sections, warnings, title=title)


# ── Plain text / Markdown / CSV ─────────────────────────────────────────
def extract_text(data: bytes, markdown: bool) -> Extraction:
    text = _decode(data).replace("\r\n", "\n")
    sections: list[Section] = []
    if markdown:
        heading, start_line, lines = "", 1, []
        def flush(end_line):
            body = clean_text("\n".join(lines))
            if body:
                loc = f"{heading} · lines {start_line}–{end_line}" if heading else f"Lines {start_line}–{end_line}"
                sections.append(Section(loc, body, heading=heading))
        for number, line in enumerate(text.split("\n"), 1):
            match = re.match(r"^(#{1,6})\s+(.*)", line)
            if match:
                flush(number - 1)
                heading, start_line, lines = _inline(match.group(2))[:80], number + 1, []
            else:
                lines.append(line)
        flush(len(text.split("\n")))
    else:
        block, start, size = [], 1, 0
        all_lines = text.split("\n")
        for number, line in enumerate(all_lines, 1):
            if not block:
                start = number
            block.append(line)
            size += len(line)
            if size > 3000 and not line.strip():
                body = clean_text("\n".join(block))
                if body:
                    sections.append(Section(f"Lines {start}–{number}", body))
                block, size = [], 0
        body = clean_text("\n".join(block))
        if body:
            sections.append(Section(f"Lines {start}–{len(all_lines)}", body))
    return Extraction("Markdown" if markdown else "Text", sections)


def extract_csv(data: bytes) -> Extraction:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return Extraction("CSV", [])
    header = [h.strip() or f"Column {i+1}" for i, h in enumerate(rows[0])]
    body = rows[1:C.MAX_CSV_ROWS + 1]
    warnings = []
    if len(rows) - 1 > C.MAX_CSV_ROWS:
        warnings.append(f"Only the first {C.MAX_CSV_ROWS:,} rows were indexed.")
    warnings.append("Each row is indexed as 'column: value' text. For totals or aggregates, use a spreadsheet tool.")
    sections = []
    for start in range(0, len(body), 20):
        group = body[start:start + 20]
        lines = ["; ".join(f"{h}: {v.strip()}" for h, v in zip(header, row) if v.strip()) for row in group]
        sections.append(Section(f"Rows {start + 1}–{start + len(group)}", "\n".join(lines)))
    return Extraction("CSV", sections, warnings)


# ── Images ──────────────────────────────────────────────────────────────
def extract_image(data: bytes, transcribe: Optional[Transcriber]) -> Extraction:
    if not transcribe:
        raise ValueError("Image reading needs a vision model.")
    text = clean_text(transcribe(data))
    return Extraction("Image", [Section("Image transcription", text)],
                      ["AI transcription may contain errors — check important text against the original."])


# ── Dispatcher ──────────────────────────────────────────────────────────
def extract_file(name: str, data: bytes, transcribe: Optional[Transcriber] = None,
                 ocr_scanned: bool = True) -> Extraction:
    suffix = Path(name).suffix.lower().lstrip(".")
    if suffix == "pdf":
        result = extract_pdf(data, transcribe, ocr_scanned)
    elif suffix == "docx":
        result = extract_docx(data)
    elif suffix == "doc":
        raise ValueError("Legacy .doc isn't supported. Save it as .docx and upload again.")
    elif suffix in C.IMAGE_TYPES:
        result = extract_image(data, transcribe)
    elif suffix in ("txt", "md"):
        result = extract_text(data, markdown=suffix == "md")
    elif suffix == "csv":
        result = extract_csv(data)
    else:
        raise ValueError(f"Unsupported file type: .{suffix}")
    if not result.sections:
        raise ValueError("No readable text was found in this file.")
    return result
